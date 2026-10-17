"""
Carga las preguntas frecuentes de Asistencia Técnica en atv.pregunta_frecuente
(scripts/005_create_pregunta_frecuente.sql) desde el Excel de docs/.

    python -m scripts.cargar_preguntas_frecuentes "docs/Preguntas_y_Respuestas_AldeaZON rellenado.xlsx"
    python -m scripts.cargar_preguntas_frecuentes <xlsx> --dry-run   # solo valida y muestra

Hoja «Preguntas», encabezados en la fila 4 (se buscan por nombre, no por
posición). Se ignoran la fila EJEMPLO y los IDs sin pregunta.

Upsert por codigo en una sola transacción. Los códigos que ya no vienen en el
Excel quedan activo = FALSE (no se borran: atv.consulta los referencia en
registros_usados). La API recalcula sus embeddings sola al ver el cambio.

Gates (no escriben nada):
  - falta una columna obligatoria, un ID repetido o una fila sin respuesta;
  - el Excel trae menos de MIN_PROPORCION_FILAS de las preguntas activas
    vigentes (archivo equivocado o a medio llenar).
"""

import argparse
import asyncio
import logging
import os
import re
from typing import Dict, List

from openpyxl import load_workbook
from sqlalchemy import func, select, update
from sqlalchemy.dialects.postgresql import insert

from app.db.session import AsyncSessionLocal
from app.models.traceability import PreguntaFrecuente

logger = logging.getLogger("atv.cargar_preguntas_frecuentes")

HOJA = "Preguntas"
FILA_ENCABEZADO = 4
MIN_PROPORCION_FILAS = 0.5

# Encabezado del Excel (prefijo, sin distinguir mayúsculas) -> columna de la tabla
COLUMNAS = {
    "id": "codigo",
    "cultivo": "cultivo",
    "categoría": "categoria",
    "pregunta": "pregunta",
    "otras formas": "variantes",
    "respuesta correcta": "respuesta",
    "dato crítico": "dato_critico",
    "respuesta peligrosa": "respuesta_evitar",
    "nivel de riesgo": "nivel_riesgo",
    "producto": "productos",
    "fuente": "fuente",
    "notas": "notas",
}
OBLIGATORIAS = ("codigo", "pregunta", "respuesta")


class ExcelInvalidoError(ValueError):
    """El Excel no tiene el formato esperado; no se carga nada."""


def separar_variantes(texto) -> List[str]:
    """'¿Sirve X? ¿Cómo lo aplico?' -> dos variantes. También separa por salto de línea o ';'."""
    if not texto:
        return []
    partes = re.split(r"[\n;]+|(?<=\?)\s+(?=¿)", str(texto))
    return [p.strip() for p in partes if p.strip()]


def _limpio(valor):
    if valor is None:
        return None
    texto = re.sub(r"\s+", " ", str(valor)).strip()
    return texto or None


def leer_excel(ruta: str) -> List[Dict]:
    hoja = load_workbook(ruta, read_only=True, data_only=True)[HOJA]
    filas = hoja.iter_rows(min_row=FILA_ENCABEZADO, values_only=True)

    encabezado = [str(c or "").strip().lower() for c in next(filas)]
    indices = {}
    for prefijo, columna in COLUMNAS.items():
        idx = next((i for i, h in enumerate(encabezado) if h.startswith(prefijo)), None)
        if idx is not None:
            indices[columna] = idx
    faltantes = [c for c in OBLIGATORIAS + ("respuesta_evitar",) if c not in indices]
    if faltantes:
        raise ExcelInvalidoError(f"Faltan columnas en la hoja '{HOJA}': {faltantes}")

    registros, vistos = [], set()
    nombre = os.path.basename(ruta)
    for fila in filas:
        valores = {col: fila[i] if i < len(fila) else None for col, i in indices.items()}
        codigo = _limpio(valores["codigo"])
        if not codigo or codigo.upper() == "EJEMPLO" or not _limpio(valores["pregunta"]):
            continue
        if codigo in vistos:
            raise ExcelInvalidoError(f"ID repetido en el Excel: {codigo}")
        vistos.add(codigo)
        registro = {col: _limpio(v) for col, v in valores.items() if col != "variantes"}
        if not registro["respuesta"]:
            raise ExcelInvalidoError(f"{codigo} no tiene «Respuesta correcta esperada»")
        registro["codigo"] = codigo
        registro["variantes"] = separar_variantes(valores.get("variantes"))
        registro["archivo_origen"] = nombre[:255]
        registros.append(registro)
    return registros


async def cargar(ruta: str, dry_run: bool = False) -> Dict[str, int]:
    registros = leer_excel(ruta)
    async with AsyncSessionLocal() as session:
        vigentes = (await session.execute(
            select(func.count()).select_from(PreguntaFrecuente).where(PreguntaFrecuente.activo.is_(True))
        )).scalar_one()
        if vigentes and len(registros) < MIN_PROPORCION_FILAS * vigentes:
            raise ExcelInvalidoError(
                f"El Excel trae {len(registros)} preguntas y hay {vigentes} activas: "
                f"menos del {MIN_PROPORCION_FILAS:.0%}, no se carga."
            )
        resumen = {"en_excel": len(registros), "activas_antes": vigentes}
        if dry_run or not registros:
            return resumen

        stmt = insert(PreguntaFrecuente).values(
            [{**r, "activo": True, "actualizado_por": "cargar_preguntas_frecuentes"} for r in registros]
        )
        columnas = [c for c in registros[0] if c != "codigo"] + ["activo", "actualizado_por"]
        stmt = stmt.on_conflict_do_update(
            index_elements=["codigo"],
            set_={**{c: stmt.excluded[c] for c in columnas}, "ts_actualizado": func.now()},
        )
        await session.execute(stmt)
        desactivadas = await session.execute(
            update(PreguntaFrecuente)
            .where(PreguntaFrecuente.codigo.not_in([r["codigo"] for r in registros]),
                   PreguntaFrecuente.activo.is_(True))
            .values(activo=False, ts_actualizado=func.now(), actualizado_por="cargar_preguntas_frecuentes")
        )
        await session.commit()
        resumen["desactivadas"] = desactivadas.rowcount
        return resumen


def main():
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("excel")
    parser.add_argument("--dry-run", action="store_true", help="valida y muestra, sin escribir")
    args = parser.parse_args()
    resumen = asyncio.run(cargar(args.excel, args.dry_run))
    logger.info(f"Preguntas frecuentes{' (dry-run)' if args.dry_run else ''}: {resumen}")


if __name__ == "__main__":
    main()
